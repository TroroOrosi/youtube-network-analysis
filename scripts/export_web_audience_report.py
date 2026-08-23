"""Build aggregate-only Web and Excel audience-network report artifacts."""

from __future__ import annotations

import json
import re
import sys
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

import analysis_common as ac

JSON_OUTPUT = REPO_ROOT / "web_ui" / "data" / "audience-network-analysis.json"
XLSX_OUTPUT = REPO_ROOT / "web_ui" / "assets" / "audience-network-analysis.xlsx"
BLUE = "0B5CAB"
PALE_BLUE = "EAF4FF"
WHITE = "FFFFFF"


def build_document() -> dict[str, Any]:
    """Return the bounded, identifier-free aggregate report document."""

    data = ac.load_all()
    audience = data["edges_anon"]
    network = data["network_edges"]
    community_summary = data["community_summary"]
    overlap = data["panel_overlap"]
    breadth = ac.viewer_breadth(audience)
    viewers = int(audience["viewer"].nunique())

    top_channels = ac.audience_top_channels(audience, top_n=50)
    categories = ac.interest_breakdown(audience, top_n_channels=300)
    communities = community_summary.sort_values("n_channels", ascending=False)
    affinity = ac.top_affinity(overlap, top_n=50, min_commenters=10)
    popularity = ac.affinity_vs_popularity(overlap, min_commenters=10)
    popularity = popularity.sort_values("rank_gap", ascending=False).head(50)

    retrieved = pd.to_datetime(overlap["retrieved_at"], utc=True)
    first_date = retrieved.min().date().isoformat()
    last_date = retrieved.max().date().isoformat()
    coverage = first_date if first_date == last_date else f"{first_date}〜{last_date}"

    return {
        "schema_version": 1,
        "title": "視聴者ネットワーク分析",
        "source_label": "匿名化済み既存調査データ（デモ）",
        "coverage_period": f"データ取得日: {coverage}",
        "methodology": (
            "公開されている購読関係を匿名集計した参考分析です。"
            "ログイン中のチャンネルから今回収集したデータではありません。"
            "非公開の購読はYouTube APIから取得できません。"
        ),
        "summary": {
            "viewers": viewers,
            "channels": int(audience["channel_id"].nunique()),
            "viewer_channel_edges": len(audience),
            "network_nodes": len(set(network["channel_id_a"]) | set(network["channel_id_b"])),
            "network_edges": len(network),
            "communities": int(community_summary["community_id"].nunique()),
            "median_breadth": int(breadth.median()),
        },
        "top_channels": [
            {
                "channel": str(row.channel_title),
                "viewers": int(row.viewers),
                "panel_share": float(row.panel_share),
            }
            for row in top_channels.itertuples(index=False)
        ],
        "viewer_breadth": _breadth_buckets(breadth, viewers),
        "categories": [
            {
                "category": str(row.category),
                "channels": int(row.n_channels),
                "viewer_links": int(row.viewer_links),
                "share": float(row.link_share),
            }
            for row in categories.itertuples(index=False)
        ],
        "communities": [
            {
                "rank": rank,
                "channels": int(row.n_channels),
                "edges": int(row.n_edges),
                "internal_weight": float(row.internal_weight),
                "representatives": str(row.representative_channels),
            }
            for rank, row in enumerate(communities.itertuples(index=False), start=1)
        ],
        "network_relationships": _network_relationships(data),
        "affinity": [
            {
                "channel": str(row.channel),
                "observed_viewers": int(row.observed_commenters),
                "panel_share": float(row.sample_share),
                "affinity_lift": float(row.affinity_lift),
            }
            for row in affinity.itertuples(index=False)
        ],
        "popularity_affinity": [
            {
                "channel": str(row.channel),
                "popularity_rank": int(row.popularity_rank),
                "affinity_rank": int(row.affinity_rank),
                "rank_gap": int(row.rank_gap),
                "panel_share": float(row.sample_share),
                "affinity_lift": float(row.affinity_lift),
            }
            for row in popularity.itertuples(index=False)
        ],
    }


def _breadth_buckets(breadth: pd.Series, viewers: int) -> list[dict[str, Any]]:
    ranges = (
        (1, 49, "1〜49"),
        (50, 99, "50〜99"),
        (100, 249, "100〜249"),
        (250, 499, "250〜499"),
        (500, 749, "500〜749"),
        (750, 1_000, "750〜1,000"),
    )
    return [
        {
            "label": label,
            "viewers": count,
            "share": count / viewers,
        }
        for lower, upper, label in ranges
        if (count := int(((breadth >= lower) & (breadth <= upper)).sum())) > 0
    ]


def _network_relationships(data: dict[str, pd.DataFrame]) -> list[dict[str, Any]]:
    audience_titles = (
        data["edges_anon"]
        .dropna(subset=["channel_title"])
        .drop_duplicates("channel_id")
        .set_index("channel_id")["channel_title"]
        .to_dict()
    )
    community_titles = (
        data["communities"]
        .dropna(subset=["channel"])
        .drop_duplicates("channel_id")
        .set_index("channel_id")["channel"]
        .to_dict()
    )
    titles = {**audience_titles, **community_titles}
    network = data["network_edges"].copy()
    network["left"] = network["channel_id_a"].map(titles)
    network["right"] = network["channel_id_b"].map(titles)
    network = network.dropna(subset=["left", "right"])
    network = network[network["left"] != network["right"]]
    network = network.sort_values(
        ["shared_viewers", "weighted_cosine"], ascending=False
    ).head(50)
    return [
        {
            "left": str(row.left),
            "right": str(row.right),
            "shared_viewers": int(row.shared_viewers),
            "weighted_cosine": float(row.weighted_cosine),
        }
        for row in network.itertuples(index=False)
    ]


def write_artifacts(
    document: dict[str, Any],
    json_path: Path = JSON_OUTPUT,
    workbook_path: Path = XLSX_OUTPUT,
) -> None:
    json_path.parent.mkdir(parents=True, exist_ok=True)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    _write_workbook(document, workbook_path)


def _write_workbook(document: dict[str, Any], output: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    fixed_time = datetime(2026, 6, 22, tzinfo=UTC).replace(tzinfo=None)
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    workbook.properties.title = document["title"]
    workbook.properties.subject = document["source_label"]

    _summary_sheet(workbook, document)
    _table_sheet(
        workbook,
        "Top Channels",
        "よく一緒に登録されているチャンネル",
        document["top_channels"],
        (("channel", "チャンネル"), ("viewers", "視聴者数"), ("panel_share", "構成比")),
        chart=True,
        percentage_columns=(3,),
    )
    _table_sheet(
        workbook,
        "Viewer Breadth",
        "視聴者ごとの登録チャンネル数",
        document["viewer_breadth"],
        (("label", "登録数"), ("viewers", "視聴者数"), ("share", "構成比")),
        chart=True,
        percentage_columns=(3,),
    )
    _table_sheet(
        workbook,
        "Interest Categories",
        "興味カテゴリ",
        document["categories"],
        (
            ("category", "カテゴリ"),
            ("channels", "チャンネル数"),
            ("viewer_links", "視聴者リンク数"),
            ("share", "構成比"),
        ),
        chart=True,
        percentage_columns=(4,),
    )
    _table_sheet(
        workbook,
        "Communities",
        "コミュニティ",
        document["communities"],
        (
            ("rank", "順位"),
            ("channels", "チャンネル数"),
            ("edges", "内部エッジ数"),
            ("internal_weight", "内部重み"),
            ("representatives", "代表チャンネル"),
        ),
    )
    _table_sheet(
        workbook,
        "Network Relationships",
        "強いネットワーク関係",
        document["network_relationships"],
        (
            ("left", "チャンネルA"),
            ("right", "チャンネルB"),
            ("shared_viewers", "共通視聴者数"),
            ("weighted_cosine", "重み付き類似度"),
        ),
        percentage_columns=(4,),
    )
    _table_sheet(
        workbook,
        "Affinity",
        "親和度",
        document["affinity"],
        (
            ("channel", "チャンネル"),
            ("observed_viewers", "観測視聴者数"),
            ("panel_share", "パネル比率"),
            ("affinity_lift", "親和度リフト"),
        ),
        percentage_columns=(3,),
    )
    _table_sheet(
        workbook,
        "Popularity vs Affinity",
        "人気度と親和度の差",
        document["popularity_affinity"],
        (
            ("channel", "チャンネル"),
            ("popularity_rank", "人気順位"),
            ("affinity_rank", "親和度順位"),
            ("rank_gap", "順位差"),
            ("panel_share", "パネル比率"),
            ("affinity_lift", "親和度リフト"),
        ),
        percentage_columns=(5,),
    )
    _methodology_sheet(workbook, document)

    with tempfile.TemporaryDirectory() as directory:
        raw_path = Path(directory) / "raw.xlsx"
        workbook.save(raw_path)
        _canonicalize_xlsx(raw_path, output)


def _summary_sheet(workbook: Workbook, document: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Summary")
    _configure_print(sheet, fit_height=1)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = document["title"]
    sheet["A1"].font = Font(size=18, bold=True, color=BLUE)
    sheet["A2"] = document["source_label"]
    sheet["A3"] = document["coverage_period"]
    labels = (
        ("viewers", "視聴者数"),
        ("channels", "登録チャンネル数"),
        ("viewer_channel_edges", "視聴者・チャンネル関係数"),
        ("network_nodes", "ネットワークノード数"),
        ("network_edges", "ネットワークエッジ数"),
        ("communities", "コミュニティ数"),
        ("median_breadth", "登録チャンネル数中央値"),
    )
    sheet.append([])
    sheet.append(["指標", "値"])
    for key, label in labels:
        sheet.append([label, document["summary"][key]])
    _style_table(sheet, 5, 5 + len(labels), 2)
    sheet["A14"] = "重要"
    sheet["A15"] = document["methodology"]
    sheet["A15"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("A15:F17")
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 18


def _table_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    rows: list[dict[str, Any]],
    columns: tuple[tuple[str, str], ...],
    *,
    chart: bool = False,
    percentage_columns: tuple[int, ...] = (),
) -> None:
    sheet = workbook.create_sheet(name)
    _configure_print(sheet, fit_height=0, repeat_header=True)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True, color=BLUE)
    sheet.append([])
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append([row[key] for key, _ in columns])
    last_row = 3 + len(rows)
    _style_table(sheet, 3, last_row, len(columns))
    for column in percentage_columns:
        for cell in sheet.iter_cols(
            min_col=column, max_col=column, min_row=4, max_row=last_row
        ):
            for item in cell:
                item.number_format = "0.0%"
    for column, (key, _) in enumerate(columns, start=1):
        width = 44 if key in {"channel", "left", "right", "representatives"} else 18
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A4"
    if chart and rows:
        chart_object = BarChart()
        chart_object.type = "bar"
        chart_object.style = 10
        chart_object.title = title
        chart_object.height = 8
        chart_object.width = 14
        data = Reference(sheet, min_col=2, min_row=3, max_row=min(last_row, 13))
        categories = Reference(sheet, min_col=1, min_row=4, max_row=min(last_row, 13))
        chart_object.add_data(data, titles_from_data=True)
        chart_object.set_categories(categories)
        chart_object.legend = None
        sheet.add_chart(chart_object, get_column_letter(len(columns) + 2) + "3")


def _style_table(sheet: Any, header_row: int, last_row: int, columns: int) -> None:
    for cell in sheet[header_row][:columns]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(
        min_row=header_row + 1, max_row=last_row, min_col=1, max_col=columns
    ):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            if isinstance(cell.value, int):
                cell.number_format = "#,##0"
            elif isinstance(cell.value, float):
                cell.number_format = "0.00"


def _methodology_sheet(workbook: Workbook, document: dict[str, Any]) -> None:
    sheet = workbook.create_sheet("Methodology")
    _configure_print(sheet, fit_height=1)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "データソースと読み方"
    sheet["A1"].font = Font(size=16, bold=True, color=BLUE)
    rows = (
        ("データ区分", document["source_label"]),
        ("取得時点", document["coverage_period"]),
        ("方法", document["methodology"]),
        ("共通購読", "同じ視聴者が登録しているチャンネルを人数と比率で集計。"),
        ("コミュニティ", "共通購読ネットワークをLouvain法でグループ化。"),
        ("親和度", "全体規模で補正したリフト。1より大きいほど固有の関連が強い。"),
        ("制約", "非公開購読は取得できず、結果は全視聴者を代表しません。"),
    )
    sheet.append([])
    sheet.append(["項目", "説明"])
    for row in rows:
        sheet.append(list(row))
    _style_table(sheet, 3, 3 + len(rows), 2)
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 88


def _configure_print(
    sheet: Any, *, fit_height: int, repeat_header: bool = False
) -> None:
    """Keep each report sheet readable when printed or saved as PDF."""

    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = fit_height
    if repeat_header:
        sheet.print_title_rows = "1:3"


def _canonicalize_xlsx(source: Path, output: Path) -> None:
    """Normalize ZIP timestamps and order so unchanged data yields no diff."""

    fixed = (1980, 1, 1, 0, 0, 0)
    with zipfile.ZipFile(source, "r") as incoming, zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as outgoing:
        for name in sorted(incoming.namelist()):
            info = zipfile.ZipInfo(name, fixed)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            payload = incoming.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-06-22T00:00:00Z\g<2>",
                    payload,
                )
            outgoing.writestr(info, payload)


def main() -> None:
    document = build_document()
    write_artifacts(document)
    print(f"wrote {JSON_OUTPUT.relative_to(REPO_ROOT)}")
    print(f"wrote {XLSX_OUTPUT.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
