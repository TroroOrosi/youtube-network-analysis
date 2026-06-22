"""Build the Excel summary workbook for the YouTube audience-network analysis.

Run:
    python scripts/generate_figures.py     # produces figures/ first
    python scripts/build_excel_report.py    # writes the workbook

Output: youtube_network_analysis_report.xlsx

Sheets
------
  Overview              project summary + headline findings
  Data Profile          per-dataset row/column/null/dup profile
  Data Dictionary       column-level descriptions
  Quality Checks        automated PASS/WARN checks
  Summary Metrics       single-number KPIs
  Top Channels          "what viewers also watch" (video's headline question)
  Affinity Lift         size-corrected audience-specific affinity
  Communities           Louvain co-viewing clusters
  Charts                embedded PNGs from figures/
  Sources & Methodology external references with URLs
"""

from __future__ import annotations

from datetime import datetime, timezone

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

import analysis_common as ac

# --- styling ---------------------------------------------------------------
NAVY = "1F3864"
BLUE = "2E5496"
LIGHT = "D9E1F2"
BAND = "EEF3FB"
GREEN = "C6EFCE"
AMBER = "FFEB9C"
GREY = "808080"

HEADER_FILL = PatternFill("solid", fgColor=BLUE)
TITLE_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor=BAND)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE_BOLD = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color="FFFFFF", size=16)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def _style_header_row(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER


def _autofit(ws, widths: dict[int, int]):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _write_df(ws, df, start_row=1, start_col=1, band=True, int_cols=(),
              pct_cols=(), float_cols=()):
    """Write a DataFrame with a styled header, banding, and number formats."""
    headers = list(df.columns)
    for j, h in enumerate(headers):
        ws.cell(row=start_row, column=start_col + j, value=h)
    _style_header_row(ws, start_row, len(headers), start_col)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        r = start_row + i
        for j, h in enumerate(headers):
            cell = ws.cell(row=r, column=start_col + j, value=row[h])
            cell.border = BORDER
            if band and i % 2 == 0:
                cell.fill = BAND_FILL
            if h in int_cols:
                cell.number_format = "#,##0"
            elif h in pct_cols:
                cell.number_format = "0.0%"
            elif h in float_cols:
                cell.number_format = "0.000"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=start_col)
    last_col = get_column_letter(start_col + len(headers) - 1)
    ws.auto_filter.ref = f"{get_column_letter(start_col)}{start_row}:{last_col}{start_row + len(df)}"
    return start_row + len(df) + 1


def _title_block(ws, title, subtitle=None):
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = title
    c.fill = TITLE_FILL
    c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34
    if subtitle:
        ws.merge_cells("A2:H2")
        s = ws["A2"]
        s.value = subtitle
        s.font = Font(italic=True, color=GREY, size=10)
        s.alignment = Alignment(horizontal="left", indent=1)


# --- sheet builders --------------------------------------------------------
def sheet_overview(wb, data, kpis):
    ws = wb.active
    ws.title = "Overview"
    _title_block(ws, "YouTube オーディエンス・ネットワーク分析レポート",
                 "「視聴者は普段どんなチャンネルを見ているのか？」を集計データで再現")
    ws.sheet_view.showGridLines = False

    rows = [
        ("生成日時", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("分析の問い",
         "自チャンネルの視聴者（コメント投稿者の代理）が、他にどんなチャンネルへ"
         "登録しているかを集計・ネットワーク化して把握する。"),
        ("参考動画",
         "「視聴者は普段どんなチャンネルを見ているのか？」"
         "（チャンネル：データで語る棒人間, 7:38）"),
        ("手法の要約",
         "コメント投稿者の公開登録チャンネルから 視聴者×チャンネル の2部グラフを"
         "構築 → 共起（共視聴）ネットワークへ射影 → Louvain でコミュニティ抽出 → "
         "規模補正した親和性 lift を算出。出力は集計のみ（個人特定なし）。"),
        ("", ""),
        ("主要な発見", ""),
        ("1. 観測パネル規模",
         f"視聴者 {kpis['n_viewers']:,} 人、ユニーク登録先 {kpis['n_channels']:,} "
         f"チャンネル、エッジ {kpis['n_edges']:,} 本。"),
        ("2. 視聴者が他に見るTop",
         f"最頻の併用登録先は「{kpis['top1_title']}」"
         f"（パネルの {kpis['top1_share']*100:.0f}% が登録）。"
         "自己啓発・ビジネス・教養系が上位を占める。"),
        ("3. コミュニティ構造",
         f"{kpis['n_communities']} クラスタを検出。最大クラスタは "
         f"{kpis['biggest_community']:,} チャンネル。"
         "「ビジネス/自己啓発」「エンタメ/音楽」など系統が分離。"),
        ("4. 固有親和性",
         f"規模補正後の lift 最大は {kpis['top_lift']:.1f}×。"
         "大型チャンネルの汎用人気と、自視聴者に固有の親和性を区別できる。"),
        ("", ""),
        ("重要な限界", ""),
        ("L1",
         "視聴者の代理は「コメント投稿者かつ公開登録者」。登録を非公開にしている"
         "層は観測できず、母数は 808 人と小さい（選択バイアスあり）。"),
        ("L2",
         "YouTube Data/Analytics API には『視聴者が他に見るチャンネル』は無い。"
         "唯一の地上検証は YouTube Studio『視聴者』タブ（オーナー権限・手動）。"),
        ("L3",
         "動画のトランスクリプトはbot判定で取得不可。動画の具体的な数値手法は"
         "公開メタデータからの推定であり、本分析は『最も近い再現』である。"),
    ]
    r = 4
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True, size=11,
                                                          color=NAVY)
        ws.cell(row=r, column=1).alignment = WRAP
        cell = ws.cell(row=r, column=2, value=val)
        cell.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        if val and len(str(val)) > 60:
            ws.row_dimensions[r].height = 44
        r += 1
    _autofit(ws, {1: 22, 2: 18, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 14})


def sheet_data_profile(wb, data):
    ws = wb.create_sheet("Data Profile")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "データプロファイル", "収集された各データセットの構造と品質指標")
    import pandas as pd
    rows = []
    for name, df in data.items():
        rows.append({
            "データセット": name,
            "行数": len(df),
            "列数": df.shape[1],
            "欠損セル合計": int(df.isna().sum().sum()),
            "完全重複行": int(df.duplicated().sum()),
            "主な列": ", ".join(map(str, df.columns[:6])),
        })
    prof = pd.DataFrame(rows)
    _write_df(ws, prof, start_row=4, int_cols=("行数", "列数", "欠損セル合計",
                                               "完全重複行"))
    _autofit(ws, {1: 32, 2: 12, 3: 8, 4: 14, 5: 12, 6: 60})


def sheet_data_dictionary(wb):
    ws = wb.create_sheet("Data Dictionary")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "データ辞書", "列ごとの意味と算出根拠")
    import pandas as pd
    rows = [
        ("edges_anon", "channel_id", "登録先チャンネルID"),
        ("edges_anon", "channel_title", "登録先チャンネル名"),
        ("edges_anon", "viewer", "匿名化された視聴者ID（u0..）。コメント投稿者の代理"),
        ("channel_network_edges_filtered", "channel_id_a/b", "共起チャンネルのペア"),
        ("channel_network_edges_filtered", "shared_viewers", "両方に登録する共通視聴者数（≥3で採用）"),
        ("channel_network_edges_filtered", "fractional_overlap", "小さい側の登録者数で正規化した重なり率"),
        ("channel_network_edges_filtered", "weighted_cosine", "次数で重み付けしたコサイン類似度（エッジ重み）"),
        ("channel_communities", "community_id", "Louvain クラスタID"),
        ("channel_communities", "sample_viewers", "このチャンネルに登録するパネル視聴者数"),
        ("channel_communities", "internal_strength", "同一クラスタ内エッジ重みの合計（ハブ度）"),
        ("channel_community_summary", "n_channels / n_edges", "クラスタの規模（チャンネル数 / 内部エッジ数）"),
        ("channel_community_summary", "internal_weight", "クラスタ内エッジ重みの総和"),
        ("channel_community_summary", "representative_channels", "内部強度上位の代表チャンネル"),
        ("channel_panel_overlap_metrics", "observed_commenters", "このチャンネルに登録する観測コメント投稿者数"),
        ("channel_panel_overlap_metrics", "observed_panel_size", "公開登録が取得できたパネル総数（808）"),
        ("channel_panel_overlap_metrics", "sample_share", "パネル内浸透率 = observed / panel_size"),
        ("channel_panel_overlap_metrics", "sample_share_lcb95", "Wilson 95% 下限（少数サンプル補正）"),
        ("channel_panel_overlap_metrics", "total_subs", "チャンネル総登録者数（Data API）"),
        ("channel_panel_overlap_metrics", "penetration_per_1M_subs", "登録者100万人あたりのパネル浸透"),
        ("channel_panel_overlap_metrics", "affinity_lift", "規模補正済み親和性（中央値=1.0、>1で過剰登録）"),
        ("channel_statistics_cache", "total_subs", "Data API から取得した総登録者数キャッシュ"),
        ("channel_statistics_cache", "subscriber_count_hidden", "登録者数が非公開か"),
    ]
    dd = pd.DataFrame(rows, columns=["データセット", "列", "説明"])
    _write_df(ws, dd, start_row=4)
    _autofit(ws, {1: 34, 2: 24, 3: 70})


def sheet_quality(wb, data):
    ws = wb.create_sheet("Quality Checks")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "品質チェック", "自動判定（PASS / WARN）")
    import pandas as pd
    ea = data["edges_anon"]
    ne = data["network_edges"]
    po = data["panel_overlap"]
    checks = [
        ("edges_anon に重複(channel_id,viewer)が無い",
         "PASS" if ea.duplicated(["channel_id", "viewer"]).sum() == 0 else "WARN",
         f"{ea.duplicated(['channel_id','viewer']).sum()} 件"),
        ("共起エッジは shared_viewers>=3",
         "PASS" if ne["shared_viewers"].min() >= 3 else "WARN",
         f"min={ne['shared_viewers'].min()}"),
        ("weighted_cosine は (0,1] の範囲",
         "PASS" if ne["weighted_cosine"].between(0, 1.0001).all() else "WARN",
         f"max={ne['weighted_cosine'].max():.3f}"),
        ("affinity_lift の中央値が約1.0に校正",
         "PASS" if abs(po["affinity_lift"].median() - 1.0) < 0.05 else "WARN",
         f"median={po['affinity_lift'].median():.3f}"),
        ("panel サイズが全行で一定(=観測パネル)",
         "PASS" if po["observed_panel_size"].nunique() == 1 else "WARN",
         f"unique={po['observed_panel_size'].nunique()}"),
        ("チャンネル統計の登録者数が非負",
         "PASS" if (data["channel_stats"]["total_subs"] >= 0).all() else "WARN",
         f"min={data['channel_stats']['total_subs'].min():,}"),
        ("観測パネル規模が小さい（解釈上の注意）",
         "WARN", f"viewers={ea['viewer'].nunique()}（選択バイアスに注意）"),
    ]
    qc = pd.DataFrame(checks, columns=["チェック項目", "判定", "詳細"])
    end = _write_df(ws, qc, start_row=4)
    # color the judgement column
    for i in range(len(qc)):
        cell = ws.cell(row=5 + i, column=2)
        cell.alignment = CENTER
        cell.fill = PatternFill("solid",
                                fgColor=GREEN if cell.value == "PASS" else AMBER)
        cell.font = Font(bold=True)
    _autofit(ws, {1: 44, 2: 10, 3: 40})


def sheet_summary_metrics(wb, kpis):
    ws = wb.create_sheet("Summary Metrics")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "サマリー指標", "主要KPI（単一値）")
    import pandas as pd
    rows = [
        ("観測視聴者（パネル）数", kpis["n_viewers"], "公開登録が取れたコメント投稿者"),
        ("ユニーク登録先チャンネル数", kpis["n_channels"], "視聴者が登録する全チャンネル"),
        ("視聴者×チャンネル エッジ数", kpis["n_edges"], "2部グラフのエッジ"),
        ("視聴者あたり登録数（中央値）", kpis["median_breadth"], "1人が登録するチャンネル数"),
        ("共起ネットワーク・エッジ数", kpis["n_net_edges"], "shared_viewers>=3"),
        ("ネットワーク・ノード数", kpis["n_net_nodes"], "共起グラフ上のチャンネル"),
        ("検出コミュニティ数", kpis["n_communities"], "Louvain クラスタ"),
        ("最大コミュニティ規模", kpis["biggest_community"], "チャンネル数"),
        ("最大親和性 lift", round(kpis["top_lift"], 2), "規模補正後・中央値=1.0"),
        ("最頻併用登録先のパネル浸透", round(kpis["top1_share"], 4), kpis["top1_title"]),
    ]
    sm = pd.DataFrame(rows, columns=["指標", "値", "説明"])
    _write_df(ws, sm, start_row=4, int_cols=("値",))
    # last two rows are not ints -> fix format
    ws.cell(row=4 + 9, column=2).number_format = "0.00"
    ws.cell(row=4 + 10, column=2).number_format = "0.0%"
    _autofit(ws, {1: 34, 2: 16, 3: 44})


def sheet_top_channels(wb, data):
    ws = wb.create_sheet("Top Channels")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "視聴者が他に見ているチャンネル Top50",
                 "動画の中心的な問いの再現：パネル視聴者の併用登録先ランキング")
    top = ac.audience_top_channels(data["edges_anon"], top_n=50)
    top.insert(0, "順位", range(1, len(top) + 1))
    top = top.rename(columns={"channel_title": "チャンネル名", "viewers": "登録視聴者数",
                              "panel_share": "パネル浸透率", "channel_id": "チャンネルID"})
    top = top[["順位", "チャンネル名", "登録視聴者数", "パネル浸透率", "チャンネルID"]]
    _write_df(ws, top, start_row=4, int_cols=("順位", "登録視聴者数"),
              pct_cols=("パネル浸透率",))
    _autofit(ws, {1: 6, 2: 42, 3: 14, 4: 14, 5: 26})


def sheet_affinity(wb, data):
    ws = wb.create_sheet("Affinity Lift")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "固有親和性 lift Top40",
                 "規模で補正した親和性（>1 = サイズの割に自視聴者へ過剰登録）")
    top = ac.top_affinity(data["panel_overlap"], top_n=40)
    cols = ["channel", "observed_commenters", "sample_share", "total_subs",
            "penetration_per_1M_subs", "affinity_lift"]
    top = top[cols].rename(columns={
        "channel": "チャンネル名", "observed_commenters": "観測登録者数",
        "sample_share": "パネル浸透率", "total_subs": "総登録者数",
        "penetration_per_1M_subs": "100万人あたり浸透", "affinity_lift": "親和性lift"})
    top.insert(0, "順位", range(1, len(top) + 1))
    _write_df(ws, top, start_row=4,
              int_cols=("順位", "観測登録者数", "総登録者数"),
              pct_cols=("パネル浸透率",),
              float_cols=("100万人あたり浸透", "親和性lift"))
    _autofit(ws, {1: 6, 2: 40, 3: 12, 4: 12, 5: 14, 6: 16, 7: 12})


def sheet_communities(wb, data):
    ws = wb.create_sheet("Communities")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "共視聴コミュニティ", "Louvain クラスタと代表チャンネル")
    cs = ac.community_sizes(data["community_summary"], min_channels=2)
    cs = cs.rename(columns={
        "community_id": "ID", "n_channels": "チャンネル数", "n_edges": "内部エッジ数",
        "internal_weight": "内部重み", "representative_channels": "代表チャンネル"})
    _write_df(ws, cs, start_row=4, int_cols=("ID", "チャンネル数", "内部エッジ数"),
              float_cols=("内部重み",))
    _autofit(ws, {1: 6, 2: 12, 3: 12, 4: 12, 5: 90})


def sheet_charts(wb):
    ws = wb.create_sheet("Charts")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "チャート", "figures/ から埋め込み（再生成: scripts/generate_figures.py）")
    images = [
        ("01_top_channels.png", "視聴者が他に見ているチャンネル Top20"),
        ("02_viewer_breadth.png", "視聴者1人あたりの登録チャンネル数"),
        ("03_community_sizes.png", "コミュニティ規模 Top12"),
        ("04_affinity_lift.png", "固有親和性 lift Top20"),
        ("05_penetration_vs_size.png", "規模 vs 浸透率 散布図"),
        ("06_edge_strength.png", "共起エッジ強度の分布"),
    ]
    row = 4
    for fname, caption in images:
        path = ac.FIGURES_DIR / fname
        if not path.exists():
            continue
        cap = ws.cell(row=row, column=1, value=caption)
        cap.font = Font(bold=True, size=12, color=NAVY)
        img = XLImage(str(path))
        # scale to a consistent width
        scale = 720 / img.width
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
        ws.add_image(img, f"A{row + 1}")
        row += int(img.height / 18) + 4
    ws.column_dimensions["A"].width = 110


def sheet_sources(wb):
    ws = wb.create_sheet("Sources & Methodology")
    ws.sheet_view.showGridLines = False
    _title_block(ws, "出典・手法", "方法論と背景の参照先（URL付き）")
    import pandas as pd
    rows = [
        ("参考動画", "視聴者は普段どんなチャンネルを見ているのか？（データで語る棒人間）",
         "https://www.youtube.com/watch?v=ec26DzgBKIU"),
        ("YouTube Data API", "channels/commentThreads/subscriptions エンドポイント仕様",
         "https://developers.google.com/youtube/v3/docs"),
        ("YouTube Analytics API", "視聴者指標（『他に見るチャンネル』は非提供）",
         "https://developers.google.com/youtube/analytics"),
        ("NetworkX", "グラフ構築・中心性・連結成分",
         "https://networkx.org/documentation/stable/"),
        ("Louvain 法", "コミュニティ検出（モジュラリティ最大化）",
         "https://python-louvain.readthedocs.io/"),
        ("Adjusted Rand Index", "クラスタ安定性の指標（scikit-learn）",
         "https://scikit-learn.org/stable/modules/generated/sklearn.metrics.adjusted_rand_score.html"),
        ("Wilson score interval", "少数サンプルの二項割合の信頼下限",
         "https://en.wikipedia.org/wiki/Binomial_proportion_confidence_interval"),
        ("Okabe-Ito palette", "色覚多様性に配慮した配色",
         "https://jfly.uni-koeln.de/color/"),
        ("Cosine similarity", "共起ベクトルの類似度（エッジ重み）",
         "https://en.wikipedia.org/wiki/Cosine_similarity"),
    ]
    src = pd.DataFrame(rows, columns=["カテゴリ", "内容", "URL"])
    end = _write_df(ws, src, start_row=4)
    # hyperlink the URL column
    for i in range(len(src)):
        cell = ws.cell(row=5 + i, column=3)
        cell.hyperlink = cell.value
        cell.font = Font(color="0563C1", underline="single")
    methodology = (
        "手法（要約）: (1) 自チャンネルのコメント投稿者を視聴者の代理として収集。"
        "(2) 各投稿者の『公開』登録チャンネルを取得し 視聴者×チャンネル の2部グラフを構築（匿名化）。"
        "(3) 共通視聴者≥3 のチャンネル対を共起エッジとし、次数で重み付けしたコサイン類似度を重みに射影。"
        "(4) Louvain でコミュニティを検出。(5) チャンネル規模で割り、中央値=1.0 に校正した親和性 lift を算出。"
        "(6) 非復元サブサンプリング＋ARI でクラスタ安定性を確認。出力は集計のみ・個人特定なし。"
    )
    mr = end + 1
    ws.cell(row=mr, column=1, value="方法論ノート").font = Font(bold=True, color=NAVY)
    mc = ws.cell(row=mr + 1, column=1, value=methodology)
    mc.alignment = WRAP
    ws.merge_cells(start_row=mr + 1, start_column=1, end_row=mr + 6, end_column=3)
    _autofit(ws, {1: 22, 2: 56, 3: 70})


def compute_kpis(data):
    ea = data["edges_anon"]
    ne = data["network_edges"]
    cs = data["community_summary"]
    po = data["panel_overlap"]
    top = ac.audience_top_channels(ea, top_n=1).iloc[0]
    nodes = set(ne["channel_id_a"]) | set(ne["channel_id_b"])
    return {
        "n_viewers": int(ea["viewer"].nunique()),
        "n_channels": int(ea["channel_id"].nunique()),
        "n_edges": int(len(ea)),
        "median_breadth": int(ac.viewer_breadth(ea).median()),
        "n_net_edges": int(len(ne)),
        "n_net_nodes": int(len(nodes)),
        "n_communities": int(cs["community_id"].nunique()),
        "biggest_community": int(cs["n_channels"].max()),
        "top_lift": float(po["affinity_lift"].max()),
        "top1_title": str(top["channel_title"]),
        "top1_share": float(top["panel_share"]),
    }


def main() -> str:
    ac.configure_fonts()
    data = ac.load_all()
    kpis = compute_kpis(data)

    wb = openpyxl.Workbook()
    sheet_overview(wb, data, kpis)
    sheet_data_profile(wb, data)
    sheet_data_dictionary(wb)
    sheet_quality(wb, data)
    sheet_summary_metrics(wb, kpis)
    sheet_top_channels(wb, data)
    sheet_affinity(wb, data)
    sheet_communities(wb, data)
    sheet_charts(wb)
    sheet_sources(wb)

    out = ac.REPO_ROOT / "youtube_network_analysis_report.xlsx"
    wb.save(out)
    print(f"wrote {out.relative_to(ac.REPO_ROOT)}  ({out.stat().st_size//1024} KB)")
    return str(out)


if __name__ == "__main__":
    main()
