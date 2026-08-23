from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts import analysis_common as ac
from scripts.export_web_audience_report import build_document, write_artifacts


class ExportWebAudienceReportTests(unittest.TestCase):
    def test_generated_snapshot_reconciles_every_analysis_family(self) -> None:
        document = build_document()
        data = ac.load_all()
        audience = data["edges_anon"]
        network = data["network_edges"]
        breadth = ac.viewer_breadth(audience)

        self.assertEqual(document["summary"]["viewers"], audience["viewer"].nunique())
        self.assertEqual(
            document["summary"]["channels"], audience["channel_id"].nunique()
        )
        self.assertEqual(document["summary"]["viewer_channel_edges"], len(audience))
        self.assertEqual(
            document["summary"]["network_nodes"],
            len(set(network["channel_id_a"]) | set(network["channel_id_b"])),
        )
        self.assertEqual(document["summary"]["network_edges"], len(network))
        self.assertEqual(
            document["summary"]["communities"],
            data["community_summary"]["community_id"].nunique(),
        )
        self.assertEqual(document["summary"]["median_breadth"], int(breadth.median()))
        for name in (
            "top_channels",
            "viewer_breadth",
            "categories",
            "communities",
            "network_relationships",
            "affinity",
            "popularity_affinity",
        ):
            self.assertGreater(len(document[name]), 0, name)

        self.assertEqual(
            len(document["top_channels"]),
            len(ac.audience_top_channels(data["edges_anon"], top_n=50)),
        )
        self.assertEqual(
            len(document["categories"]),
            len(ac.interest_breakdown(data["edges_anon"], top_n_channels=300)),
        )
        self.assertEqual(
            len(document["communities"]), len(data["community_summary"])
        )
        self.assertEqual(
            len(document["affinity"]),
            len(ac.top_affinity(data["panel_overlap"], top_n=50, min_commenters=10)),
        )
        breadth_ranges = (
            (1, 49),
            (50, 99),
            (100, 249),
            (250, 499),
            (500, 749),
            (750, 1_000),
        )
        self.assertEqual(
            len(document["viewer_breadth"]),
            sum(
                ((breadth >= lower) & (breadth <= upper)).any()
                for lower, upper in breadth_ranges
            ),
        )
        popularity = ac.affinity_vs_popularity(
            data["panel_overlap"], min_commenters=10
        )
        self.assertEqual(len(document["popularity_affinity"]), min(50, len(popularity)))

        title_ids = set(
            data["edges_anon"].dropna(subset=["channel_title"])["channel_id"]
        ) | set(data["communities"].dropna(subset=["channel"])["channel_id"])
        eligible_relationships = data["network_edges"][
            data["network_edges"]["channel_id_a"].isin(title_ids)
            & data["network_edges"]["channel_id_b"].isin(title_ids)
        ]
        self.assertEqual(
            len(document["network_relationships"]), min(50, len(eligible_relationships))
        )

        serialized = json.dumps(document, ensure_ascii=False).lower()
        for forbidden in (
            '"viewer":',
            '"viewer_ids":',
            '"channel_id":',
            "viewer_001",
            "c:\\\\",
            "/users/",
            "token",
        ):
            self.assertNotIn(forbidden, serialized)

    def test_web_and_excel_artifacts_share_headline_metrics(self) -> None:
        document = build_document()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            json_path = root / "report.json"
            workbook_path = root / "report.xlsx"

            write_artifacts(document, json_path, workbook_path)
            workbook = load_workbook(workbook_path, read_only=False, data_only=False)
            calculated_workbook = load_workbook(
                workbook_path, read_only=False, data_only=True
            )

        self.assertEqual(workbook.sheetnames, [
            "Summary",
            "Top Channels",
            "Viewer Breadth",
            "Interest Categories",
            "Communities",
            "Network Relationships",
            "Affinity",
            "Popularity vs Affinity",
            "Methodology",
        ])
        summary = workbook["Summary"]
        metrics = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(6, 13)}
        self.assertEqual(metrics["視聴者数"], document["summary"]["viewers"])
        self.assertEqual(metrics["登録チャンネル数"], document["summary"]["channels"])
        self.assertEqual(len(workbook["Top Channels"]._charts), 1)
        self.assertEqual(len(workbook["Viewer Breadth"]._charts), 1)
        self.assertEqual(len(workbook["Interest Categories"]._charts), 1)

        summary = workbook["Summary"]
        self.assertTrue(summary.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(summary.page_setup.fitToWidth, 1)
        self.assertEqual(summary.page_setup.fitToHeight, 1)

        for name in workbook.sheetnames[1:-1]:
            sheet = workbook[name]
            self.assertEqual(sheet.page_setup.orientation, "landscape", name)
            self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage, name)
            self.assertEqual(sheet.page_setup.fitToWidth, 1, name)
            self.assertEqual(sheet.page_setup.fitToHeight, 0, name)
            self.assertEqual(sheet.print_title_rows, "$1:$3", name)

        methodology = workbook["Methodology"]
        self.assertEqual(methodology.page_setup.orientation, "landscape")
        self.assertTrue(methodology.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(methodology.page_setup.fitToWidth, 1)
        self.assertEqual(methodology.page_setup.fitToHeight, 1)

        formula_cells = [
            cell.coordinate
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        error_values = [
            cell.value
            for sheet in calculated_workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "e"
            or cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
        ]
        self.assertEqual(formula_cells, [])
        self.assertEqual(error_values, [])
        workbook_text = "\n".join(
            str(cell.value).lower()
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for forbidden in (
            "viewer_001",
            "channel_id",
            "c:\\users\\",
            "/users/",
            "token",
        ):
            self.assertNotIn(forbidden, workbook_text)

    def test_generated_input_is_validated_before_either_artifact_is_written(self) -> None:
        document = build_document()
        document["viewer_ids"] = ["viewer_001"]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            json_path = root / "report.json"
            workbook_path = root / "report.xlsx"

            with self.assertRaisesRegex(ValueError, "unexpected report fields"):
                write_artifacts(document, json_path, workbook_path)

            self.assertFalse(json_path.exists())
            self.assertFalse(workbook_path.exists())

    def test_formula_leading_text_is_inert_in_excel(self) -> None:
        document = build_document()
        document["top_channels"][0]["channel"] = "=1+1"
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "report.xlsx"

            write_artifacts(document, root / "report.json", workbook_path)
            workbook = load_workbook(workbook_path, data_only=False)

        cell = workbook["Top Channels"]["A4"]
        self.assertEqual(cell.value, "'=1+1")
        self.assertNotEqual(cell.data_type, "f")

    def test_empty_analysis_families_still_generate_a_readable_workbook(self) -> None:
        document = build_document()
        for name in (
            "top_channels",
            "viewer_breadth",
            "categories",
            "communities",
            "network_relationships",
            "affinity",
            "popularity_affinity",
        ):
            document[name] = []
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "report.xlsx"

            write_artifacts(document, root / "report.json", workbook_path)
            workbook = load_workbook(workbook_path, data_only=False)

        self.assertEqual(len(workbook.sheetnames), 9)
        self.assertEqual(
            sum(len(sheet._charts) for sheet in workbook.worksheets), 0
        )


if __name__ == "__main__":
    unittest.main()
